from __future__ import annotations

from pathlib import Path

import pytest
from pptx import Presentation

from pdf2pptx.cli import main
from tests.pdfgen import make_pdf


def test_default_output_next_to_input(beamer_pdf: Path, capsys: pytest.CaptureFixture[str]) -> None:
    assert main([str(beamer_pdf)]) == 0
    out = beamer_pdf.with_suffix(".pptx")
    assert len(Presentation(str(out)).slides) == 4
    assert "4 slides from 4 pages" in capsys.readouterr().out


def test_explicit_output_file(beamer_pdf: Path, tmp_path: Path) -> None:
    out = tmp_path / "sub" / "deck.pptx"
    assert main([str(beamer_pdf), "-o", str(out), "--collapse-overlays", "-q"]) == 0
    assert len(Presentation(str(out)).slides) == 3


def test_refuses_to_overwrite_without_force(
    beamer_pdf: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    out = beamer_pdf.with_suffix(".pptx")
    out.write_bytes(b"precious")
    assert main([str(beamer_pdf)]) == 1
    assert "already exists" in capsys.readouterr().err
    assert out.read_bytes() == b"precious"
    assert main([str(beamer_pdf), "--force"]) == 0
    assert out.read_bytes() != b"precious"


def test_many_inputs_into_directory(tmp_path: Path) -> None:
    a = make_pdf(tmp_path / "a.pdf", ["x"])
    b = make_pdf(tmp_path / "b.pdf", ["y", "z"])
    outdir = tmp_path / "out"
    outdir.mkdir()
    assert main([str(a), str(b), "-o", str(outdir), "-q"]) == 0
    assert len(Presentation(str(outdir / "a.pptx")).slides) == 1
    assert len(Presentation(str(outdir / "b.pptx")).slides) == 2


def test_many_inputs_rejects_pptx_output(tmp_path: Path) -> None:
    a = make_pdf(tmp_path / "a.pdf", ["x"])
    assert main([str(a), str(a), "-o", str(tmp_path / "x.pptx")]) == 2


def test_one_bad_input_fails_run_but_converts_the_rest(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    good = make_pdf(tmp_path / "good.pdf", ["ok"])
    assert main([str(tmp_path / "missing.pdf"), str(good), "-q"]) == 1
    assert "input not found" in capsys.readouterr().err
    assert good.with_suffix(".pptx").exists()


def test_output_must_be_pptx(beamer_pdf: Path, tmp_path: Path) -> None:
    assert main([str(beamer_pdf), "-o", str(tmp_path / "deck.ppt")]) == 1


def test_bad_env_config_fails_fast(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("PDF2PPTX_DPI", "lots")
    with pytest.raises(SystemExit, match="PDF2PPTX_DPI must be an integer"):
        main(["x.pdf"])


def test_env_dpi_out_of_range_is_an_error(
    beamer_pdf: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setenv("PDF2PPTX_DPI", "5")
    assert main([str(beamer_pdf)]) == 1

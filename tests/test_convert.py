from __future__ import annotations

import io
from pathlib import Path

import pypdfium2 as pdfium
import pytest
from PIL import Image
from pptx import Presentation
from pptx.presentation import Presentation as Deck
from pptx.shapes.picture import Picture
from pptx.util import Emu

from pdf2pptx.convert import (
    EMU_PER_INCH,
    ConversionError,
    Options,
    convert,
    fit,
    slide_size_in,
)
from tests.pdfgen import make_pdf


def _pictures(prs: Deck) -> list[Picture]:
    return [s for slide in prs.slides for s in slide.shapes if isinstance(s, Picture)]


def _width_in(prs: Deck) -> float:
    assert prs.slide_width is not None
    return float(prs.slide_width) / EMU_PER_INCH


def test_one_slide_per_page_with_full_bleed_image(beamer_pdf: Path, tmp_path: Path) -> None:
    out = tmp_path / "talk.pptx"
    result = convert(beamer_pdf, out)

    assert result.pdf_pages == 4
    assert result.slides == 4
    prs = Presentation(str(out))
    assert len(prs.slides) == 4
    pics = _pictures(prs)
    assert len(pics) == 4
    assert prs.slide_width is not None and prs.slide_height is not None
    for pic in pics:
        assert (pic.left, pic.top) == (0, 0)
        assert abs(pic.width - prs.slide_width) <= 1
        assert abs(pic.height - prs.slide_height) <= 1


def test_beamer_4_3_maps_to_10_by_7_5_inches(beamer_pdf: Path, tmp_path: Path) -> None:
    out = convert(beamer_pdf, tmp_path / "t.pptx").output
    prs = Presentation(str(out))
    assert prs.slide_height is not None
    assert _width_in(prs) == pytest.approx(10.0, abs=0.01)
    assert prs.slide_height / EMU_PER_INCH == pytest.approx(7.5, abs=0.01)


def test_16_9_deck_gets_widescreen_slides(tmp_path: Path) -> None:
    pdf = make_pdf(tmp_path / "wide.pdf", ["a"], size=(1600, 900))
    prs = Presentation(str(convert(pdf, tmp_path / "w.pptx").output))
    assert _width_in(prs) == pytest.approx(13.333, abs=0.01)


def test_image_resolution_follows_dpi(beamer_pdf: Path, tmp_path: Path) -> None:
    out = convert(beamer_pdf, tmp_path / "t.pptx", Options(dpi=100)).output
    pic = _pictures(Presentation(str(out)))[0]
    img = Image.open(io.BytesIO(pic.image.blob))
    assert img.width == pytest.approx(1000, abs=2)  # 10 in * 100 dpi
    assert pic.image.content_type == "image/png"


def test_jpeg_format(beamer_pdf: Path, tmp_path: Path) -> None:
    out = convert(beamer_pdf, tmp_path / "t.pptx", Options(image_format="jpeg")).output
    assert _pictures(Presentation(str(out)))[0].image.content_type == "image/jpeg"


def test_rendered_image_contains_page_content(beamer_pdf: Path, tmp_path: Path) -> None:
    out = convert(beamer_pdf, tmp_path / "t.pptx", Options(dpi=72)).output
    img = Image.open(io.BytesIO(_pictures(Presentation(str(out)))[0].image.blob)).convert("L")
    lo, hi = img.getextrema()
    assert isinstance(lo, int) and isinstance(hi, int)
    assert hi > 200 and lo < 80, "slide image should have dark text on a light page"


def test_text_goes_to_notes_and_alt_text(beamer_pdf: Path, tmp_path: Path) -> None:
    prs = Presentation(str(convert(beamer_pdf, tmp_path / "t.pptx").output))
    slide = prs.slides[2]
    assert slide.notes_slide.notes_text_frame.text == "Results E=mc2"
    pic = _pictures(prs)[2]
    assert pic._element.xpath("./p:nvPicPr/p:cNvPr/@descr") == ["Results E=mc2"]


def test_no_notes_option(beamer_pdf: Path, tmp_path: Path) -> None:
    prs = Presentation(str(convert(beamer_pdf, tmp_path / "t.pptx", Options(notes=False)).output))
    assert all(not s.has_notes_slide for s in prs.slides)


def test_collapse_overlays_keeps_last_step(beamer_pdf: Path, tmp_path: Path) -> None:
    opts = Options(collapse_overlays=True)
    result = convert(beamer_pdf, tmp_path / "t.pptx", opts)
    assert (result.pdf_pages, result.slides) == (4, 3)
    notes = [s.notes_slide.notes_text_frame.text for s in Presentation(str(result.output)).slides]
    assert notes == ["Intro step two", "Results E=mc2", "Thanks"]


def test_collapse_overlays_without_labels_keeps_everything(tmp_path: Path) -> None:
    pdf = make_pdf(tmp_path / "plain.pdf", ["a", "b", "c"])
    result = convert(pdf, tmp_path / "t.pptx", Options(collapse_overlays=True))
    assert result.slides == 3


def test_mixed_page_sizes_are_letterboxed(tmp_path: Path) -> None:
    pdf = make_pdf(tmp_path / "mixed.pdf", ["wide", "tall"], sizes=[(400, 300), (300, 400)])
    prs = Presentation(str(convert(pdf, tmp_path / "t.pptx").output))
    tall = _pictures(prs)[1]
    assert tall.height == prs.slide_height
    assert tall.width < prs.slide_width
    assert tall.left == (prs.slide_width - tall.width) // 2


@pytest.mark.parametrize(
    ("content", "message"),
    [
        (b"hello, not a pdf", "not a PDF"),
        (b"%PDF-1.7\ngarbage that is not a real document", "could not open PDF"),
    ],
)
def test_rejects_bad_input(tmp_path: Path, content: bytes, message: str) -> None:
    bad = tmp_path / "bad.pdf"
    bad.write_bytes(content)
    with pytest.raises(ConversionError, match=message):
        convert(bad, tmp_path / "out.pptx")
    assert not (tmp_path / "out.pptx").exists()
    assert list(tmp_path.glob(".*partial")) == []


def test_missing_input(tmp_path: Path) -> None:
    with pytest.raises(ConversionError, match="input not found"):
        convert(tmp_path / "nope.pdf", tmp_path / "out.pptx")


def test_password_protected_pdf(
    beamer_pdf: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    def locked(*_a: object, **_k: object) -> None:
        raise pdfium.PdfiumError("Failed to load document (PDFium: Incorrect password error).")

    monkeypatch.setattr(pdfium, "PdfDocument", locked)
    with pytest.raises(ConversionError, match="password-protected"):
        convert(beamer_pdf, tmp_path / "out.pptx")


@pytest.mark.parametrize("opts", [Options(dpi=10), Options(dpi=5000), Options(jpeg_quality=0)])
def test_invalid_options(beamer_pdf: Path, tmp_path: Path, opts: Options) -> None:
    with pytest.raises(ConversionError):
        convert(beamer_pdf, tmp_path / "out.pptx", opts)


def test_slide_size_clamped_to_powerpoint_limits() -> None:
    w, h = slide_size_in(10_000, 10)  # absurd banner
    assert w == 56.0 and 1.0 <= h <= 56.0


def test_fit_centres() -> None:
    assert fit(1, 1, 200, 100) == (50, 0, 100, 100)
    assert isinstance(Emu(fit(4, 3, 400, 300)[2]), int)

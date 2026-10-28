"""Core conversion: render every PDF page to an image and place it on its own slide."""

from __future__ import annotations

import io
from collections.abc import Iterator
from dataclasses import dataclass
from pathlib import Path
from typing import Literal

import pypdfium2 as pdfium
from pptx import Presentation
from pptx.util import Emu

ImageFormat = Literal["png", "jpeg"]

EMU_PER_INCH = 914_400
# PowerPoint's default slide height; width follows the PDF's aspect ratio
# (4:3 beamer -> 10 x 7.5 in, 16:9 -> 13.33 x 7.5 in).
SLIDE_HEIGHT_IN = 7.5
# PowerPoint rejects slide dimensions outside 1-56 inches.
MIN_SLIDE_IN, MAX_SLIDE_IN = 1.0, 56.0
MIN_DPI, MAX_DPI = 50, 600
PDF_MAGIC = b"%PDF-"


class ConversionError(Exception):
    """Raised for bad input that the user can fix (not a bug in this tool)."""


@dataclass(frozen=True)
class Options:
    dpi: int = 220
    image_format: ImageFormat = "png"
    jpeg_quality: int = 90
    notes: bool = True
    collapse_overlays: bool = False

    def validate(self) -> None:
        if not MIN_DPI <= self.dpi <= MAX_DPI:
            raise ConversionError(f"dpi must be between {MIN_DPI} and {MAX_DPI}, got {self.dpi}")
        if self.image_format not in ("png", "jpeg"):
            raise ConversionError(f"image format must be png or jpeg, got {self.image_format!r}")
        if not 1 <= self.jpeg_quality <= 100:
            raise ConversionError(f"jpeg quality must be 1-100, got {self.jpeg_quality}")


@dataclass(frozen=True)
class Result:
    output: Path
    pdf_pages: int
    slides: int


def open_pdf(path: Path) -> pdfium.PdfDocument:
    if not path.is_file():
        raise ConversionError(f"input not found: {path}")
    with path.open("rb") as fh:
        head = fh.read(1024)
    # The spec allows junk before the header; 1 KiB is what most readers tolerate.
    if PDF_MAGIC not in head:
        raise ConversionError(f"not a PDF file: {path}")
    try:
        doc = pdfium.PdfDocument(str(path))
    except pdfium.PdfiumError as exc:
        msg = str(exc)
        if "password" in msg.lower():
            raise ConversionError(f"PDF is password-protected: {path}") from exc
        raise ConversionError(f"could not open PDF ({msg}): {path}") from exc
    if len(doc) == 0:
        doc.close()
        raise ConversionError(f"PDF has no pages: {path}")
    return doc


def select_pages(doc: pdfium.PdfDocument, collapse_overlays: bool) -> list[int]:
    """Return the page indices to convert.

    Beamer emits one PDF page per overlay step (\\pause, \\only, ...) and gives
    all steps of a frame the same page label. Collapsing keeps only the last,
    fully built-up step of each run of identically labelled pages.
    """
    n = len(doc)
    if not collapse_overlays:
        return list(range(n))
    labels = [doc.get_page_label(i) for i in range(n)]
    if not any(labels):
        return list(range(n))
    keep = []
    for i in range(n):
        last_in_run = i == n - 1 or not labels[i] or labels[i] != labels[i + 1]
        if last_in_run:
            keep.append(i)
    return keep


def slide_size_in(page_w_pt: float, page_h_pt: float) -> tuple[float, float]:
    if page_w_pt <= 0 or page_h_pt <= 0:
        raise ConversionError("PDF page has zero size")
    aspect = page_w_pt / page_h_pt
    height = SLIDE_HEIGHT_IN
    width = height * aspect
    if width > MAX_SLIDE_IN:
        width, height = MAX_SLIDE_IN, MAX_SLIDE_IN / aspect
    elif width < MIN_SLIDE_IN:
        width, height = MIN_SLIDE_IN, MIN_SLIDE_IN / aspect
    height = min(max(height, MIN_SLIDE_IN), MAX_SLIDE_IN)
    return width, height


def fit(img_w: float, img_h: float, box_w: int, box_h: int) -> tuple[int, int, int, int]:
    """Largest (left, top, width, height) of img aspect centred inside the box."""
    scale = min(box_w / img_w, box_h / img_h)
    w, h = round(img_w * scale), round(img_h * scale)
    return (box_w - w) // 2, (box_h - h) // 2, w, h


def render_page(page: pdfium.PdfPage, target_w_in: float, opts: Options) -> bytes:
    page_w_pt, _ = page.get_size()
    # Render so the image has `dpi` pixels per inch of the *slide*, not of the
    # PDF page: beamer pages are only ~5 in wide, so page DPI would be blurry.
    scale = (target_w_in * opts.dpi) / page_w_pt
    image = page.render(scale=scale, may_draw_forms=True).to_pil()
    buf = io.BytesIO()
    if opts.image_format == "jpeg":
        image.convert("RGB").save(buf, format="JPEG", quality=opts.jpeg_quality, optimize=True)
    else:
        image.save(buf, format="PNG", optimize=True)
    return buf.getvalue()


def page_text(page: pdfium.PdfPage) -> str:
    textpage = page.get_textpage()
    try:
        text: str = textpage.get_text_bounded()
    finally:
        textpage.close()
    lines = (line.strip() for line in text.replace("\r\n", "\n").replace("\r", "\n").split("\n"))
    return "\n".join(line for line in lines if line)


def _iter_pages(doc: pdfium.PdfDocument, indices: list[int]) -> Iterator[pdfium.PdfPage]:
    for i in indices:
        page = doc[i]
        try:
            yield page
        finally:
            page.close()


def convert(src: Path, dst: Path, opts: Options | None = None) -> Result:
    opts = opts or Options()
    opts.validate()
    doc = open_pdf(src)
    try:
        indices = select_pages(doc, opts.collapse_overlays)
        first_w, first_h = doc.get_page_size(indices[0])
        slide_w_in, slide_h_in = slide_size_in(first_w, first_h)

        prs = Presentation()
        slide_w = Emu(round(slide_w_in * EMU_PER_INCH))
        slide_h = Emu(round(slide_h_in * EMU_PER_INCH))
        prs.slide_width, prs.slide_height = slide_w, slide_h
        blank = prs.slide_layouts[6]

        for n, page in enumerate(_iter_pages(doc, indices), start=1):
            pw, ph = page.get_size()
            left, top, w, h = fit(pw, ph, slide_w, slide_h)
            data = render_page(page, w / EMU_PER_INCH, opts)
            slide = prs.slides.add_slide(blank)
            pic = slide.shapes.add_picture(io.BytesIO(data), left, top, w, h)
            text = page_text(page)
            pic.name = f"Slide {n}"
            # Alt text so screen readers get something better than "Picture 1".
            pic._element.nvPicPr.cNvPr.set("descr", text[:1000] or f"Slide {n}")
            if opts.notes and text:
                slide.notes_slide.notes_text_frame.text = text

        dst.parent.mkdir(parents=True, exist_ok=True)
        tmp = dst.with_name(f".{dst.name}.partial")
        try:
            prs.save(str(tmp))
            tmp.replace(dst)  # atomic: never leave a half-written deck behind
        finally:
            tmp.unlink(missing_ok=True)
        return Result(output=dst, pdf_pages=len(doc), slides=len(indices))
    finally:
        doc.close()
